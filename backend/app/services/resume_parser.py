import io, re
import fitz
from docx import Document

def extract_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            with fitz.open(stream=data, filetype="pdf") as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        elif lower.endswith(".docx"):
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
        else:
            raise ValueError("Only PDF and DOCX files are supported.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read document: {exc}") from exc
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 30:
        raise ValueError("No readable text found. This may be a scanned PDF; OCR is not enabled.")
    return text
